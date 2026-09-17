from pathlib import Path
import re

import chromadb
from openpyxl import load_workbook
from pypdf import PdfReader


# Resolve paths from the repository root so the script works from any current directory.
ROOT_DIR = Path(__file__).resolve().parent.parent
PDF_DIR = ROOT_DIR / "doc"
CHROMA_DIR = ROOT_DIR / "data" / "chroma"
METADATA_PATH = PDF_DIR / "policy_company_metadata.xlsx"
COLLECTION_NAME = "health_insurance_policies"


# Convert spreadsheet headers into stable Chroma metadata keys.
def normalize_metadata_key(header):
    return re.sub(r"[^a-z0-9]+", "_", str(header).lower()).strip("_")


# Read the workbook's company rows and index them by their PDF filename.
def load_policy_metadata():
    workbook = load_workbook(METADATA_PATH, read_only=True, data_only=True)
    worksheet = workbook["Company Metadata"]
    rows = list(worksheet.iter_rows(values_only=True))
    header_index = next(
        index
        for index, row in enumerate(rows)
        if row and row[0] == "Company Name"
    )
    headers = [normalize_metadata_key(value) for value in rows[header_index]]
    metadata_by_file = {}

    for row in rows[header_index + 1 :]:
        if not row or not row[0]:
            continue
        metadata = {
            key: value
            for key, value in zip(headers, row)
            if key and value is not None
        }
        filename = metadata.get("policy_document_file")
        if filename:
            metadata_by_file[filename] = metadata

    if not metadata_by_file:
        raise ValueError(f"No policy metadata rows found in {METADATA_PATH}")
    return metadata_by_file


# Extract one searchable document per PDF page and attach its company metadata.
def extract_policy_chunks():
    """Extract page documents with metadata suitable for Chroma filtering."""
    documents = []
    metadatas = []
    ids = []
    metadata_by_file = load_policy_metadata()

    policy_files = sorted(PDF_DIR.glob("*.pdf"))
    if not policy_files:
        raise FileNotFoundError(f"No policy PDFs found in {PDF_DIR}")

    for pdf_path in policy_files:
        company_metadata = metadata_by_file.get(pdf_path.name, {})
        company_name = company_metadata.get("company_name", pdf_path.stem)
        reader = PdfReader(str(pdf_path))
        for page_number, page in enumerate(reader.pages, start=1):
            text = " ".join((page.extract_text() or "").split())
            if not text:
                continue
            ids.append(f"{pdf_path.stem.lower()}-page-{page_number}")
            documents.append(text)
            page_metadata = dict(company_metadata)
            page_metadata.update(
                {"company": company_name, "source": pdf_path.name, "page": page_number}
            )
            metadatas.append(page_metadata)

    if not documents:
        raise ValueError("No policy text was extracted from the insurer PDFs")
    return ids, documents, metadatas


# Upsert all policy page embeddings into the persistent Chroma collection.
def build_database():
    """Create or refresh the persistent ChromaDB collection with workbook metadata."""
    ids, documents, metadatas = extract_policy_chunks()
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    # Replace the prior collection so removed or renamed source files cannot leave stale vectors.
    try:
        client.delete_collection(COLLECTION_NAME)
    except ValueError:
        pass
    collection = client.get_or_create_collection(
        name=COLLECTION_NAME,
        metadata={"description": "Vector embeddings for health insurance policies"},
    )
    collection.upsert(ids=ids, documents=documents, metadatas=metadatas)
    print(f"Collection: {collection.name}")
    print(f"Stored policy page embeddings: {collection.count()}")
    print(f"Database path: {CHROMA_DIR}")
    return collection


# Search the collection and optionally restrict results by any metadata field.
def query_policies(query_text, where=None, n_results=5):
    """Return nearest policy pages, optionally filtered by Chroma metadata."""
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    collection = client.get_collection(COLLECTION_NAME)
    return collection.query(query_texts=[query_text], where=where, n_results=n_results)


# Build the database when this file is executed as a script.
if __name__ == "__main__":
    build_database()